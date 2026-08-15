from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable


STRUCTURED_INPUT_KINDS = {
    "three_list_tasks",
    "three_list_issues",
    "work_logs",
}

_SCHEMA_FIELDS = {
    "three_list_tasks": {"任务描述", "任务状态"},
    "three_list_issues": {"问题描述", "状态"},
    "work_logs": {"工作内容", "姓名", "日期"},
}


class StructuredWorkbookError(ValueError):
    pass


@dataclass(frozen=True)
class WorkbookRow:
    sheet_name: str
    row_number: int
    values: dict[str, Any]


def resolve_project_input_kind(filename: str, content: bytes, requested: str = "auto") -> str:
    normalized = str(requested or "auto").strip().lower()
    suffix = Path(filename).suffix.lower()
    if normalized in STRUCTURED_INPUT_KINDS:
        detected = detect_structured_workbook_kind(filename, content)
        if detected != normalized:
            raise StructuredWorkbookError(
                f"上传类型 {normalized} 与 Excel 表头识别结果 {detected} 不一致。"
            )
        return normalized
    if normalized == "minutes":
        if suffix not in {".md", ".markdown", ".txt", ".docx"}:
            raise ValueError(f"不支持的会议纪要格式：{suffix}，请使用 .md、.txt 或 .docx")
        return normalized
    if normalized == "transcript":
        if suffix not in {".txt", ".docx"}:
            raise ValueError(f"不支持的转写格式：{suffix}，请使用 .txt 或 .docx")
        return normalized
    if normalized != "auto":
        raise ValueError(f"不支持的 input_kind：{requested}")
    if suffix in {".xlsx", ".xlsm"}:
        return detect_structured_workbook_kind(filename, content)
    if suffix in {".md", ".markdown", ".txt"}:
        return "minutes"
    if suffix == ".docx":
        return "transcript"
    raise ValueError(
        f"无法自动识别 {suffix or '无扩展名文件'}；支持会议纪要、原始转写、任务清单、问题清单和工作日志。"
    )


def detect_structured_workbook_kind(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise StructuredWorkbookError(f"结构化表格只支持 .xlsx 或 .xlsm，当前为 {suffix or '无扩展名'}。")
    workbook = _load_workbook(BytesIO(content), filename)
    try:
        matches: set[str] = set()
        observed_headers: list[str] = []
        for sheet in workbook.worksheets:
            headers = _sheet_headers(sheet)
            if not headers:
                continue
            observed_headers.extend(headers)
            matches.update(_matching_kinds(headers))
    finally:
        workbook.close()
    if len(matches) == 1:
        return next(iter(matches))
    if len(matches) > 1:
        raise StructuredWorkbookError(
            "一个 Excel 同时包含多种可导入表结构；请分别导出任务、问题或工作日志后再上传。"
        )
    preview = "、".join(dict.fromkeys(observed_headers[:12])) or "空表"
    raise StructuredWorkbookError(f"无法识别 Excel 表头：{preview}")


def read_structured_workbook(path: str | Path, expected_kind: str) -> list[WorkbookRow]:
    if expected_kind not in STRUCTURED_INPUT_KINDS:
        raise StructuredWorkbookError(f"不支持的结构化表格类型：{expected_kind}")
    source_path = Path(path)
    workbook = _load_workbook(source_path, source_path.name)
    records: list[WorkbookRow] = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                raw_headers = next(rows)
            except StopIteration:
                continue
            headers = [_text(value) for value in raw_headers]
            if expected_kind not in _matching_kinds(headers):
                continue
            for row_number, values in enumerate(rows, start=2):
                if not any(value not in (None, "") for value in values):
                    continue
                records.append(
                    WorkbookRow(
                        sheet_name=sheet.title,
                        row_number=row_number,
                        values=dict(zip(headers, values)),
                    )
                )
    finally:
        workbook.close()
    if not records:
        raise StructuredWorkbookError(f"{expected_kind} 表中没有可导入的数据行。")
    return records


def _load_workbook(source: Any, filename: str) -> Any:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel 导入需要 openpyxl，请先安装 requirements.txt。") from exc
    try:
        return load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise StructuredWorkbookError(f"无法读取 Excel 文件 {filename}：{exc}") from exc


def _sheet_headers(sheet: Any) -> list[str]:
    for row in sheet.iter_rows(values_only=True):
        headers = [_text(value) for value in row]
        if any(headers):
            return headers
    return []


def _matching_kinds(headers: Iterable[str]) -> set[str]:
    header_set = {_text(value) for value in headers if _text(value)}
    return {
        kind
        for kind, required in _SCHEMA_FIELDS.items()
        if required.issubset(header_set)
    }


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()
