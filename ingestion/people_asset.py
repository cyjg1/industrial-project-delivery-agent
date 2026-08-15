from __future__ import annotations

import csv
import json
import os
import re
from pathlib import Path
from typing import Any

from agent.repository_paths import REPOSITORY_ROOT, normalize_repository_paths
from agent.schemas import PeopleGroup, PeopleStructure, Person, PersonAssignment, ScenarioOrg, to_plain


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_PEOPLE_ASSET_PATH = REPO_ROOT / "data" / "assets" / "people_structure.json"


def configured_people_asset_path(asset_path: str | Path | None = None) -> Path:
    if asset_path:
        return Path(asset_path)
    configured = os.getenv("PROJECT_AGENT_PEOPLE_ASSET_PATH", "").strip()
    return Path(configured) if configured else DEFAULT_PEOPLE_ASSET_PATH


def load_people_asset(asset_path: str | Path | None = None) -> PeopleStructure:
    path = configured_people_asset_path(asset_path)
    if not path.exists():
        raise FileNotFoundError(
            f"People asset not found: {path}. Create it from Excel, XMind, CSV, or manual JSON edits."
        )
    payload = json.loads(path.read_text(encoding="utf-8"))
    return people_structure_from_payload(payload)


def save_people_asset(
    structure: PeopleStructure,
    asset_path: str | Path | None = None,
    *,
    asset_id: str = "people_structure",
    source_type: str = "manual_update",
    source_path: str = "",
    notes: str = "",
) -> None:
    path = configured_people_asset_path(asset_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "asset_id": asset_id,
        "source": {
            "type": source_type,
            "path": source_path,
            "status": "canonical_people_asset",
            "accepted_update_channels": ["excel", "xmind", "csv", "manual"],
            "notes": notes,
        },
        "people_structure": to_plain(structure),
    }
    try:
        path.resolve().relative_to(REPOSITORY_ROOT.resolve())
        portability_root = REPOSITORY_ROOT
    except ValueError:
        portability_root = path.parent
    portable_payload = normalize_repository_paths(payload, root=portability_root)
    path.write_text(json.dumps(portable_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def people_structure_from_payload(payload: dict[str, Any]) -> PeopleStructure:
    data = payload.get("people_structure", payload)
    return PeopleStructure(
        root_title=data.get("root_title", ""),
        groups=[PeopleGroup(**item) for item in data.get("groups", [])],
        people=[_person_from_payload(item) for item in data.get("people", [])],
        scenarios=[ScenarioOrg(**item) for item in data.get("scenarios", [])],
    )


def _person_from_payload(item: dict[str, Any]) -> Person:
    payload = dict(item)
    payload["assignments"] = [
        assignment if isinstance(assignment, PersonAssignment) else PersonAssignment(**assignment)
        for assignment in payload.get("assignments", [])
    ]
    return Person(**payload)


def refresh_people_asset_from_table(
    table_path: str | Path,
    asset_path: str | Path | None = None,
) -> PeopleStructure:
    source = Path(table_path)
    structure = load_people_structure_from_table(source)
    save_people_asset(
        structure,
        asset_path=asset_path,
        source_type=_table_source_type(source),
        source_path=str(source),
        notes="Imported into the canonical people asset for meeting understanding, responsibility matching, and task assignment.",
    )
    return structure


def load_people_structure_from_table(table_path: str | Path) -> PeopleStructure:
    path = Path(table_path)
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return people_structure_from_rows(list(csv.DictReader(handle)))
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return people_structure_from_rows(_read_xlsx_rows(path))
    raise ValueError(f"Unsupported people asset table format: {path.suffix}. Use .csv, .xlsx, or .xlsm.")


def people_structure_from_rows(rows: list[dict[str, Any]]) -> PeopleStructure:
    normalized_rows = [_normalize_row(row) for row in rows]
    root_title = _first(
        _value(row, "root_title")
        for row in normalized_rows
    ) or "人员资产"

    groups: list[PeopleGroup] = []
    people: list[Person] = []
    scenarios: list[ScenarioOrg] = []
    seen_groups: set[tuple[str, str]] = set()
    seen_people: set[tuple[str, str, str]] = set()
    seen_scenarios: set[str] = set()

    for row in normalized_rows:
        record_type = _record_type(row)
        name = _value(row, "name")
        group = _value(row, "group")
        role = _value(row, "role")
        path = _value(row, "path")
        scenario_id = _value(row, "scenario_id")

        if record_type == "group" and name:
            key = (name, path or name)
            if key not in seen_groups:
                seen_groups.add(key)
                groups.append(PeopleGroup(name=name, path=path or name))
            continue

        if record_type == "person" and name:
            key = (name, group, role)
            if key not in seen_people:
                seen_people.add(key)
                people.append(
                    Person(
                        name=name,
                        group=group,
                        role=role,
                        path=path or " / ".join(item for item in [root_title, group, role, name] if item),
                        responsibility_note=_value(row, "responsibility_note"),
                        person_id=_value(row, "person_id"),
                    )
                )
            continue

        if record_type == "scenario" and scenario_id:
            if scenario_id not in seen_scenarios:
                seen_scenarios.add(scenario_id)
                scenarios.append(
                    ScenarioOrg(
                        scenario_id=scenario_id,
                        name=name,
                        teams=_split_list(_value(row, "teams") or group),
                        leads=_split_list(_value(row, "leads")),
                        path=path or " / ".join(item for item in [root_title, scenario_id, name] if item),
                    )
                )

    return PeopleStructure(
        root_title=root_title,
        groups=groups,
        people=people,
        scenarios=sorted(scenarios, key=lambda item: item.scenario_id),
    )


def _read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel people asset import requires openpyxl. Install requirements.txt first.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    return [
        {
            headers[index]: cell
            for index, cell in enumerate(row)
            if index < len(headers) and headers[index]
        }
        for row in rows[1:]
    ]


def _normalize_row(row: dict[str, Any]) -> dict[str, str]:
    return {
        str(key).strip().lower(): "" if value is None else str(value).strip()
        for key, value in row.items()
    }


def _value(row: dict[str, str], field: str) -> str:
    aliases = {
        "record_type": ("record_type", "type", "kind", "类型"),
        "root_title": ("root_title", "资产名称", "根节点"),
        "name": ("name", "名称", "姓名", "场景名称"),
        "group": ("group", "分组", "组", "团队", "所属业务板块"),
        "role": ("role", "角色", "职责"),
        "responsibility_note": ("responsibility_note", "责任说明", "负责什么", "职责说明", "职责备注"),
        "person_id": ("person_id", "人员id", "用户id", "成员id"),
        "path": ("path", "路径"),
        "scenario_id": ("scenario_id", "scenario", "场景id", "场景编号"),
        "teams": ("teams", "team", "涉及团队"),
        "leads": ("leads", "lead", "负责人", "主导人员"),
    }
    for alias in aliases[field]:
        value = row.get(alias.lower(), "")
        if value:
            return value
    return ""


def _record_type(row: dict[str, str]) -> str:
    explicit = _value(row, "record_type").lower()
    if explicit in {"group", "people_group", "组", "分组"}:
        return "group"
    if explicit in {"person", "people", "人", "人员"}:
        return "person"
    if explicit in {"scenario", "scenario_org", "uat", "场景"}:
        return "scenario"
    if _value(row, "scenario_id"):
        return "scenario"
    if _value(row, "role") or _value(row, "group"):
        return "person"
    return "group"


def _split_list(value: str) -> list[str]:
    return [
        item.strip()
        for item in re.split(r"[,;；、/\n]+", value)
        if item.strip()
    ]


def _first(values: Any) -> str:
    for value in values:
        if value:
            return value
    return ""


def _table_source_type(path: Path) -> str:
    if path.suffix.lower() == ".csv":
        return "csv_import"
    return "excel_import"
