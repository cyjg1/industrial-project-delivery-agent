from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import threading
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any

from agent.access_policy import User
from agent.meeting_ingestion import enqueue_source_ingestion
from agent.schemas import SourceDocument, SourceRef
from ingestion.source_manifest import register_uploaded_project_file
from skills.meeting_minutes.transcript import parse_transcript_bytes
from store.sqlite_store import ProjectSQLiteStore


_MANIFEST_LOCK = threading.RLock()
_TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".csv", ".json", ".yaml", ".yml"}
_WORKBOOK_SUFFIXES = {".xlsx", ".xlsm"}
_MINUTES_SUFFIXES = {".md", ".markdown", ".txt", ".docx"}
_CATEGORY_RULES = (
    ("需求资料", ("需求", "requirement", "用户故事", "原型")),
    ("设计方案", ("设计", "架构", "方案", "design", "architecture", "蓝图")),
    ("项目计划", ("计划", "排期", "进度", "里程碑", "plan", "schedule")),
    ("项目报告", ("报告", "汇报", "周报", "月报", "总结", "复盘", "report")),
    ("交付物", ("交付", "验收", "成果", "deliverable", "acceptance")),
    ("测试资料", ("测试", "uat", "用例", "缺陷", "test", "bug")),
    ("管理制度", ("制度", "规范", "流程", "标准", "policy", "procedure")),
    ("参考资料", ("参考", "手册", "指南", "说明", "reference", "manual", "guide")),
)


class ProjectFileError(ValueError):
    pass


class ProjectFileService:
    def __init__(self, store: ProjectSQLiteStore):
        self.store = store
        self.root = store.root_dir / "project_files"
        self.manifest_path = self.root / "manifest.json"
        self.index_root = self.root / ".index"

    def workspace(self, *, project_id: str) -> dict[str, Any]:
        project_root = self._project_root(project_id)
        project_root.mkdir(parents=True, exist_ok=True)
        entries = [row for row in self._read_manifest() if row.get("project_id") == project_id]
        files_by_path = {str(row.get("relative_path") or ""): row for row in entries}
        folders = {""}
        files: list[dict[str, Any]] = []
        for path in sorted(project_root.rglob("*"), key=lambda item: item.as_posix().lower()):
            if self.index_root in path.parents or path == self.index_root:
                continue
            relative = path.relative_to(project_root).as_posix()
            if path.is_dir():
                folders.add(relative)
                continue
            metadata = files_by_path.get(relative, {})
            files.append({
                "file_id": metadata.get("file_id") or self._file_id(project_id, relative),
                "name": path.name,
                "folder": (
                    "" if PurePosixPath(relative).parent.as_posix() == "."
                    else PurePosixPath(relative).parent.as_posix()
                ),
                "relative_path": relative,
                "size": path.stat().st_size,
                "modified_at": datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(),
                "upload_type": metadata.get("upload_type") or "project_document",
                "document_category": metadata.get("document_category") or "未分类",
                "source_id": metadata.get("source_id") or "",
                "job_id": metadata.get("job_id") or "",
                "index_status": metadata.get("index_status") or "未建立索引",
                "uploaded_by": metadata.get("uploaded_by") or "",
                "created_at": metadata.get("created_at") or "",
            })
        return {
            "folders": sorted(folders, key=lambda value: (value.count("/"), value.lower())),
            "files": files,
            "summary": {
                "folder_count": max(0, len(folders) - 1),
                "file_count": len(files),
                "minutes_count": sum(item["upload_type"] == "minutes" for item in files),
                "indexed_count": sum(item["index_status"] in {"已建立索引", "处理中"} for item in files),
            },
        }

    def create_folder(self, *, project_id: str, parent: str, name: str) -> str:
        clean_name = self._safe_segment(name)
        relative = self._safe_relative("/".join(filter(None, (self._safe_folder(parent), clean_name))))
        target = self._project_root(project_id) / Path(relative)
        target.mkdir(parents=True, exist_ok=True)
        return relative

    def upload(
        self,
        *,
        project_id: str,
        actor: User,
        folder: str,
        upload_type: str,
        filename: str,
        content: bytes,
        title: str = "",
        meeting_date: str = "",
        topic: str = "",
        sensitivity: str = "l1",
    ) -> tuple[dict[str, Any], dict[str, Any] | None]:
        if upload_type not in {"minutes", "project_document"}:
            raise ProjectFileError("上传类型必须是会议纪要或项目资料。")
        if not content:
            raise ProjectFileError("上传文件不能为空。")
        safe_name = self._safe_filename(filename)
        relative_folder = self._safe_folder(folder)
        target_dir = self._project_root(project_id) / Path(relative_folder)
        target_dir.mkdir(parents=True, exist_ok=True)
        target = self._unique_target(target_dir, safe_name)
        target.write_bytes(content)
        relative_path = target.relative_to(self._project_root(project_id)).as_posix()
        tags = {
            "org_id": actor.org_id,
            "project_id": project_id,
            "topic_id": None,
            "author_id": actor.id,
            "sensitivity": sensitivity,
        }
        job: dict[str, Any] | None = None
        if upload_type == "minutes":
            if target.suffix.lower() not in _MINUTES_SUFFIXES:
                target.unlink(missing_ok=True)
                raise ProjectFileError("会议纪要仅支持 .md、.txt、.markdown 或 .docx。")
            source = register_uploaded_project_file(
                filename=target.name,
                content=content,
                title=title or target.stem,
                meeting_date=meeting_date,
                topic=topic or "项目会议纪要",
                access_tags=tags,
                input_kind="minutes",
            )
            self.store.ingest(source, tags=tags, actor=actor, kind="minutes", materialize=False)
            job = enqueue_source_ingestion(store=self.store, source=source, actor=actor)
            category = "会议纪要"
            index_status = "处理中"
        else:
            text, extraction_status = self._extract_text(content, target.name)
            category = self._classify(target.name, text)
            source = self._save_project_document_source(
                project_id=project_id,
                actor=actor,
                original_path=target,
                relative_path=relative_path,
                content=content,
                text=text,
                title=title or target.stem,
                category=category,
                sensitivity=sensitivity,
            )
            index_status = "已建立索引" if extraction_status == "text_extracted" else "仅文件名索引"
        row = {
            "file_id": self._file_id(project_id, relative_path),
            "project_id": project_id,
            "name": target.name,
            "folder": relative_folder,
            "relative_path": relative_path,
            "upload_type": upload_type,
            "document_category": category,
            "source_id": source.doc_id,
            "job_id": str((job or {}).get("id") or ""),
            "index_status": index_status,
            "uploaded_by": actor.name,
            "created_at": _now(),
        }
        self._upsert_manifest(row)
        return row, job

    def resolve_download(self, *, project_id: str, relative_path: str) -> Path:
        safe_relative = self._safe_relative(relative_path)
        path = (self._project_root(project_id) / Path(safe_relative)).resolve()
        project_root = self._project_root(project_id).resolve()
        if project_root not in path.parents or not path.is_file():
            raise FileNotFoundError(relative_path)
        return path

    def _save_project_document_source(
        self,
        *,
        project_id: str,
        actor: User,
        original_path: Path,
        relative_path: str,
        content: bytes,
        text: str,
        title: str,
        category: str,
        sensitivity: str,
    ) -> SourceDocument:
        content_hash = hashlib.sha256(content).hexdigest()
        source_id = "project_doc_" + hashlib.sha256(
            f"{project_id}\0{relative_path}\0{content_hash}".encode("utf-8")
        ).hexdigest()[:24]
        index_dir = self.index_root / project_id
        index_dir.mkdir(parents=True, exist_ok=True)
        index_path = index_dir / f"{source_id}.md"
        body = text.strip() or f"# {title}\n\n文件类型：{category}\n文件名：{original_path.name}\n"
        index_path.write_text(
            f"# {title}\n\n- 资料类型：{category}\n- 原始文件：{relative_path}\n\n{body}\n",
            encoding="utf-8",
        )
        source = SourceDocument(
            doc_id=source_id,
            title=title,
            meeting_date="",
            topic=category,
            curated_source=SourceRef(
                source_type="Indexed project document",
                path=str(index_path),
                status="matched",
                notes="Searchable text projection of the uploaded project document.",
            ),
            raw_source=SourceRef(
                source_type="Uploaded project document",
                path=str(original_path),
                status="matched",
                notes="Original file retained in the project file workspace.",
            ),
            tags=["用户上传", "项目资料", category],
            input_kind="project_document",
            content_hash=content_hash,
            org_id=actor.org_id,
            project_id=project_id,
            author_id=actor.id,
            sensitivity=sensitivity,
            tag_origin="actor_ingest",
        )
        self.store.save_source(source, kind="project_document", materialize=False)
        return source

    def _extract_text(self, content: bytes, filename: str) -> tuple[str, str]:
        suffix = Path(filename).suffix.lower()
        try:
            if suffix in _TEXT_SUFFIXES:
                return _decode(content), "text_extracted"
            if suffix == ".docx":
                return parse_transcript_bytes(content, filename), "text_extracted"
            if suffix in _WORKBOOK_SUFFIXES:
                from openpyxl import load_workbook

                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                lines: list[str] = []
                for sheet in workbook.worksheets:
                    lines.append(f"## {sheet.title}")
                    for row in sheet.iter_rows(values_only=True):
                        values = [str(value).strip() for value in row if value not in (None, "")]
                        if values:
                            lines.append(" | ".join(values))
                return "\n".join(lines), "text_extracted"
        except Exception:
            return "", "metadata_only"
        return "", "metadata_only"

    def _classify(self, filename: str, text: str) -> str:
        haystack = f"{filename}\n{text[:12000]}".lower()
        scores = [
            (sum(haystack.count(keyword.lower()) for keyword in keywords), category)
            for category, keywords in _CATEGORY_RULES
        ]
        score, category = max(scores, default=(0, "其他资料"))
        return category if score else "其他资料"

    def _project_root(self, project_id: str) -> Path:
        safe_project = re.sub(r"[^A-Za-z0-9_.-]+", "_", project_id).strip("._") or "project"
        return self.root / safe_project

    def _safe_folder(self, value: str) -> str:
        value = str(value or "").strip().replace("\\", "/").strip("/")
        return self._safe_relative(value) if value else ""

    def _safe_relative(self, value: str) -> str:
        path = PurePosixPath(str(value or "").replace("\\", "/"))
        if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
            raise ProjectFileError("文件夹路径不合法。")
        return path.as_posix()

    def _safe_segment(self, value: str) -> str:
        name = str(value or "").strip()
        if not name or name in {".", ".."} or re.search(r'[<>:"/\\|?*]', name):
            raise ProjectFileError("文件夹名称不能为空，且不能包含路径或系统保留字符。")
        return name

    def _safe_filename(self, value: str) -> str:
        name = Path(str(value or "").replace("\\", "/")).name.strip()
        if not name or name in {".", ".."}:
            raise ProjectFileError("文件名不合法。")
        return re.sub(r'[<>:"/\\|?*]+', "_", name)

    def _unique_target(self, directory: Path, filename: str) -> Path:
        candidate = directory / filename
        if not candidate.exists():
            return candidate
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        return directory / f"{candidate.stem}-{stamp}{candidate.suffix}"

    def _file_id(self, project_id: str, relative_path: str) -> str:
        return "file_" + hashlib.sha256(f"{project_id}\0{relative_path}".encode()).hexdigest()[:20]

    def _read_manifest(self) -> list[dict[str, Any]]:
        with _MANIFEST_LOCK:
            if not self.manifest_path.is_file():
                return []
            try:
                value = json.loads(self.manifest_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                return []
            return value if isinstance(value, list) else []

    def _upsert_manifest(self, row: dict[str, Any]) -> None:
        with _MANIFEST_LOCK:
            rows = self._read_manifest()
            rows = [item for item in rows if item.get("file_id") != row["file_id"]]
            rows.append(row)
            self.root.mkdir(parents=True, exist_ok=True)
            temporary = self.manifest_path.with_suffix(".tmp")
            temporary.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
            shutil.move(str(temporary), str(self.manifest_path))


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()
