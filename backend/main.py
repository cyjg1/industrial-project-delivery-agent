from __future__ import annotations

import json
import os
import re
from io import BytesIO
from hashlib import sha256
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4
from zoneinfo import ZoneInfo

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from pydantic import BaseModel, Field

from agent.access_policy import IdentityResolutionError, ROLE_RANK, resolve_actor, visible
from agent.conversation.run_manager import ConversationRunManager
from agent.daily_brief import create_daily_brief_scheduler, ensure_today_daily_brief, generate_daily_brief, latest_daily_brief
from agent.daily_journal import DailyJournalScheduler, DailyJournalService, DailyJournalValidationError, compile_project_journals
from agent.deliverables import (
    DeliverableService,
    DeliverableUpload,
    DeliverableValidationError,
    MAX_DELIVERABLE_FILES,
    deliverable_file_size_limit,
    deliverable_total_size_limit,
)
from agent.evolution import build_meeting_followup_comparison, generate_weekly_review, mark_meeting_continuation
from agent.env_loader import load_project_env
from agent.file_analysis_workspace import FileAnalysisWorkspace
from agent.ingestion_worker import IngestionWorker
from agent.inspection_jobs import InspectionDispatcher, InspectionWorker
from agent.inspection_scheduler import InspectionScheduler
from agent.meeting_ingestion import MeetingIngestionError, enqueue_source_ingestion, source_ingestion_view
from agent.memory_embedding_worker import MemoryEmbeddingWorker
from agent.memory_skill_evolution import MemorySkillEvolutionService
from agent.presentation import build_human_review_workspace, ingestion_job_view
from agent.project_skills import ProjectSkillService, ProjectSkillValidationError
from agent.project_context import list_project_contexts, load_project_context, update_project_context
from agent.project_people import load_project_people, project_people_asset_path
from agent.runtime import AgentRuntime
from agent.schemas import CandidateStatus, EvidenceRef, InspectionItem, PeopleGroup, Person, to_plain
from agent.tools import build_default_registry
from agent.tool_registry import conversation_tool_specs, register_specs
from ingestion.people_asset import (
    configured_people_asset_path,
    load_people_structure_from_table,
    save_people_asset,
)
from ingestion.source_manifest import build_project_manifest, register_uploaded_project_file, upload_root
from ingestion.structured_workbook import resolve_project_input_kind
from ingestion.xmind_loader import load_people_structure_from_xmind
from store.sqlite_store import ProjectSQLiteStore

from backend.access_scope import (
    request_scope,
    require_item_edit,
    require_review,
    require_role,
    require_task_edit,
    require_task_status_edit,
    require_visible,
    resolve_request_scope,
)
from backend.routers.conversation import (
    ConversationRouterServices,
    create_conversation_router,
)
from backend.routers.memory_skill import (
    MemorySkillRouterServices,
    create_memory_skill_router,
)
from backend.routers.skills import SkillRouterServices, create_skills_router
from backend.routers.project_files import ProjectFileRouterServices, create_project_files_router
from backend.routers.runtime_config import create_runtime_config_router


DEFAULT_STORE_DIR = Path(__file__).resolve().parents[1] / "data" / "store"


class ConfirmationRequest(BaseModel):
    editor: str = "local_user"
    notes: str = ""


class EditItemRequest(ConfirmationRequest):
    status: str | None = None
    title: str | None = None
    description: str | None = None
    due_date: str | None = None
    deliverable: str | None = None
    acceptance_criteria: str | None = None
    owner_candidates: list[str] | None = None
    linked_issue_ids: list[str] | None = None
    linked_task_ids: list[str] | None = None


class MilestoneConfigRequest(BaseModel):
    milestone_id: str | None = None
    project: str | None = None
    name: str | None = None
    date_start: str | None = None
    date_end: str | None = None
    scenario_id: str | None = None
    chain_name: str | None = None
    acceptance_criteria: list[str] | str | None = None
    status: str | None = None


class WorkItemEditRequest(ConfirmationRequest):
    title: str | None = None
    description: str | None = None
    status: str | None = None
    due_date: str | None = None
    deliverable: str | None = None
    acceptance_criteria: str | None = None
    professional_id: str | None = None
    board_id: str | None = None
    planned_start: str | None = None
    progress_percent: int | None = None
    owner_candidates: list[str] | None = None
    linked_issue_ids: list[str] | None = None
    linked_task_ids: list[str] | None = None


class PersonProfileRequest(ConfirmationRequest):
    name: str | None = None
    group: str | None = None
    role: str | None = None
    path: str | None = None
    responsibility_note: str | None = None


class PeopleGroupRequest(ConfirmationRequest):
    name: str
    path: str | None = None


class BriefGenerateRequest(BaseModel):
    push_feishu: bool = False


class WorkspaceArchiveRequest(BaseModel):
    confirmation: str


class DailyJournalGenerateRequest(BaseModel):
    report_date: str | None = None


class DailyWorkRecordEditRequest(BaseModel):
    record_date: str | None = None
    kind: str | None = None
    text: str | None = None
    solution_options: list[dict[str, Any]] | None = None


class DeliverableRequirementRequest(BaseModel):
    title: str
    type_label: str = ""
    acceptance_criteria: str = ""
    due_date: str = ""
    required: bool = True
    topic_id: str | None = None
    sensitivity: str = "l1"
    sort_order: int = 0


class DeliverableRequirementEditRequest(BaseModel):
    title: str | None = None
    type_label: str | None = None
    acceptance_criteria: str | None = None
    due_date: str | None = None
    required: bool | None = None
    topic_id: str | None = None
    sensitivity: str | None = None
    sort_order: int | None = None
    status: str | None = None


class ProjectSkillCompileRequest(BaseModel):
    method_id: str


class ThreeListMutationRequest(ConfirmationRequest):
    title: str | None = None
    description: str | None = None
    status: str | None = None
    owner_candidates: list[str] | None = None
    due_date: str | None = None
    deliverable: str | None = None
    acceptance_criteria: str | None = None
    professional_id: str | None = None
    board_id: str | None = None
    linked_issue_ids: list[str] | None = None
    linked_task_ids: list[str] | None = None
    business_goal: str | None = None
    principles: list[str] | None = None
    reasoning_chain: list[str] | None = None
    applicable_scope: str | None = None


class MeetingContinuationRequest(BaseModel):
    previous_source_id: str
    next_source_id: str
    editor: str = "local_user"
    notes: str = ""


class WeeklyReviewGenerateRequest(BaseModel):
    week_end: str | None = None


class InspectionRequest(BaseModel):
    source_ids: list[str] = Field(default_factory=list)
    reason: str = ""
    request_id: str = ""


def create_app(store_dir: str | Path | None = None) -> FastAPI:
    load_project_env()
    startup_hooks: list[Any] = []
    shutdown_hooks: list[Any] = []

    @asynccontextmanager
    async def lifespan(_app: FastAPI) -> AsyncIterator[None]:
        for hook in startup_hooks:
            hook()
        try:
            yield
        finally:
            for hook in shutdown_hooks:
                hook()

    app = FastAPI(title="Project Delivery Agent API", version="0.5.0", lifespan=lifespan)
    app.add_middleware(
        CORSMiddleware,
        allow_origins=["*"],
        allow_credentials=True,
        allow_methods=["*"],
        allow_headers=["*"],
    )

    root_dir = Path(store_dir or os.getenv("PROJECT_AGENT_STORE_DIR", "") or DEFAULT_STORE_DIR)

    def get_store() -> ProjectSQLiteStore:
        return ProjectSQLiteStore(str(root_dir))

    def get_runtime(store: ProjectSQLiteStore) -> AgentRuntime:
        return AgentRuntime(tool_registry=build_default_registry(store=store), store=store)

    def get_skill_service(store: ProjectSQLiteStore) -> ProjectSkillService:
        registry = register_specs(build_default_registry(store=store), conversation_tool_specs(object()))
        return ProjectSkillService(store, registry)

    def get_file_workspace() -> FileAnalysisWorkspace:
        return FileAnalysisWorkspace()

    def get_deliverable_service(store: ProjectSQLiteStore | None = None) -> DeliverableService:
        return DeliverableService(store or get_store())

    def required_actor(request: Request, store: ProjectSQLiteStore) -> tuple[Any, Any]:
        scope = request_scope(request)
        return scope.actor, scope.access_context

    @app.middleware("http")
    async def actor_scope_middleware(request: Request, call_next: Any) -> Any:
        if (
            request.method == "OPTIONS"
            or request.url.path in {"/api/health", "/api/identity/actor"}
            or request.url.path.startswith("/api/runtime/provider-config")
            or not request.url.path.startswith("/api/")
        ):
            return await call_next(request)
        try:
            request.state.actor_scope = resolve_request_scope(request, get_store())
        except IdentityResolutionError as exc:
            return JSONResponse(status_code=401, content={"detail": str(exc)})
        except HTTPException as exc:
            return JSONResponse(status_code=exc.status_code, content={"detail": exc.detail})
        return await call_next(request)

    def latest_workspace(request: Request) -> dict[str, Any]:
        store = get_store()
        actor, access_context = required_actor(request, store)
        project_id = request_scope(request).project_id
        milestone = load_project_context(store, project_id).default_milestone
        if ROLE_RANK.get(request_scope(request).role, -1) >= ROLE_RANK["pmo"]:
            ensure_today_daily_brief(store, milestone, project_id=project_id)
        run = store.latest_run(project_id=project_id)
        payload = build_human_review_workspace(
            run=run,
            store=store,
            manifest=build_project_manifest(
                store,
                project_id,
                date_start=milestone.date_start,
                date_end=milestone.date_end,
            ),
            milestone=milestone,
            project_id=project_id,
            actor=actor,
            access_context=access_context,
        )
        payload["daily_brief"]["scheduler"] = app.state.brief_scheduler.status()
        payload["daily_journal"]["scheduler"] = app.state.journal_scheduler.status()
        return payload

    def attach_scheduler_status(payload: dict[str, Any]) -> dict[str, Any]:
        if isinstance(payload.get("workspace"), dict):
            payload["workspace"].setdefault("daily_brief", {})["scheduler"] = app.state.brief_scheduler.status()
            payload["workspace"].setdefault("daily_journal", {})["scheduler"] = app.state.journal_scheduler.status()
        return payload

    app.include_router(create_runtime_config_router())

    @app.get("/api/project-calendar")
    def project_calendar(request: Request, month: str = "") -> dict[str, Any]:
        """Return project meetings and task deadlines as a simple month calendar."""
        scope = request_scope(request)
        try:
            month_start = datetime.strptime(month or date.today().strftime("%Y-%m"), "%Y-%m").date()
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="month 必须使用 YYYY-MM 格式") from exc
        next_month = (
            date(month_start.year + 1, 1, 1)
            if month_start.month == 12
            else date(month_start.year, month_start.month + 1, 1)
        )
        store = get_store()
        actor = scope.actor
        project_id = scope.project_id
        events: list[dict[str, Any]] = []
        meeting_kinds = {"minutes", "transcript", "meeting"}

        for stored_event in store.list_calendar_events(project_id=project_id):
            try:
                event_date = datetime.strptime(str(stored_event.get("date") or ""), "%Y-%m-%d").date()
            except ValueError:
                continue
            if not month_start <= event_date < next_month:
                continue
            owner_names = [str(name) for name in stored_event.get("owner_names") or [] if str(name)]
            events.append({
                "event_id": str(stored_event.get("event_id") or ""),
                "event_type": "meeting",
                "date": event_date.isoformat(),
                "title": str(stored_event.get("title") or "项目会议"),
                "time_label": str(stored_event.get("time_label") or "会议"),
                "owner_names": owner_names,
                "status": str(stored_event.get("status") or "scheduled"),
                "deliverable": str(stored_event.get("deliverable") or ""),
                "description": str(stored_event.get("description") or ""),
                "is_related": actor.id in owner_names or actor.name in owner_names,
            })

        for source in store.list_sources():
            source_kind = str(source.get("kind") or "")
            if source.get("project_id") != project_id:
                continue
            if source_kind not in meeting_kinds:
                continue
            meeting_date = str(source.get("meeting_date") or "")
            try:
                event_date = datetime.strptime(meeting_date[:10], "%Y-%m-%d").date()
            except ValueError:
                continue
            if not month_start <= event_date < next_month:
                continue
            payload = source.get("payload") if isinstance(source.get("payload"), dict) else {}
            topic_id = str(source.get("topic_id") or payload.get("topic_id") or "")
            searchable = " ".join([
                str(source.get("title") or ""),
                str(payload.get("topic") or ""),
                " ".join(str(tag) for tag in (payload.get("tags") or [])),
            ])
            topic_members = scope.access_context.members_of(topic_id) if topic_id else set()
            is_related = (
                not topic_id
                or actor.id in topic_members
                or str(source.get("author_id") or payload.get("author_id") or "") == actor.id
                or actor.name in searchable
            )
            events.append({
                "event_id": f"meeting:{source.get('id') or source.get('doc_id')}",
                "event_type": "meeting",
                "date": event_date.isoformat(),
                "title": str(source.get("title") or "项目会议").strip("*# ") or "项目会议",
                "time_label": "会议",
                "owner_names": [],
                "status": "scheduled",
                "deliverable": "",
                "description": str(payload.get("topic") or "项目会议"),
                "is_related": is_related,
            })

        excluded_statuses = {"archived", "canceled"}
        for candidate in store.list_items(status=CandidateStatus.CANDIDATE):
            if candidate.category not in {"task", "tasks", "followup"}:
                continue
            if candidate.project_id != project_id or not candidate.due_date:
                continue
            if not visible(actor, candidate, scope.access_context):
                continue
            try:
                event_date = datetime.strptime(str(candidate.due_date)[:10], "%Y-%m-%d").date()
            except ValueError:
                continue
            if not month_start <= event_date < next_month:
                continue
            owner_names = sorted({str(name) for name in candidate.owner_candidates or [] if name})
            events.append({
                "event_id": f"task-candidate:{candidate.item_id}",
                "event_type": "task_deadline",
                "date": event_date.isoformat(),
                "title": candidate.title,
                "time_label": "截止",
                "owner_names": owner_names,
                "status": candidate.status.value,
                "deliverable": candidate.deliverable or "",
                "description": candidate.description or "",
                "is_related": actor.id in owner_names or actor.name in owner_names,
            })

        for task in store.list_work_items():
            if task.project_id != project_id or task.status.value in excluded_statuses or not task.due_date:
                continue
            try:
                event_date = datetime.strptime(str(task.due_date)[:10], "%Y-%m-%d").date()
            except ValueError:
                continue
            if not month_start <= event_date < next_month:
                continue
            related_names = {
                str(name)
                for name in [
                    *(task.owner_candidates or []),
                    *(task.collaborators or []),
                    *(task.confirmers or []),
                ]
                if name
            }
            display_title = re.sub(
                r"^(?:20)?\d{2}[-./年]?\d{2}[-./月]?\d{2}日?\s*[-—_:：|]?\s*",
                "",
                task.title.strip(),
            ) or task.title
            meeting_title = re.sub(r"[（(][^）)]*[）)]\s*$", "", display_title).strip()
            is_meeting_task = bool(re.search(
                r"(?:会议|例会|站会|沟通会|专题会|评审会|协调会|复盘会|宣贯会|启动会|碰头会)$",
                meeting_title,
            ))
            events.append({
                "event_id": f"{'meeting-task' if is_meeting_task else 'task'}:{task.work_item_id}",
                "event_type": "meeting" if is_meeting_task else "task_deadline",
                "date": event_date.isoformat(),
                "title": display_title if is_meeting_task else task.title,
                "time_label": "会议" if is_meeting_task else "截止",
                "owner_names": sorted(related_names),
                "status": task.status.value,
                "deliverable": task.deliverable or "",
                "description": task.description or "",
                "is_related": actor.id in related_names or actor.name in related_names,
            })

        events.sort(key=lambda item: (item["date"], item["event_type"] != "meeting", item["title"]))
        return {
            "month": month_start.strftime("%Y-%m"),
            "actor": {"id": actor.id, "name": actor.name},
            "events": events,
            "summary": {
                "meeting_count": sum(event["event_type"] == "meeting" for event in events),
                "deadline_count": sum(event["event_type"] == "task_deadline" for event in events),
                "related_count": sum(bool(event["is_related"]) for event in events),
            },
        }

    def scheduled_daily_brief() -> None:
        store = get_store()
        for context in list_project_contexts(store):
            generate_daily_brief(
                store=store,
                milestone=context.default_milestone,
                project_id=context.project_id,
                push_feishu=bool(os.getenv("FEISHU_WEBHOOK_URL", "")),
                webhook_url=os.getenv("FEISHU_WEBHOOK_URL", ""),
            )

    app.state.brief_scheduler = create_daily_brief_scheduler(scheduled_daily_brief)
    ingestion_worker_enabled = os.getenv("INGESTION_WORKER_ENABLED", "on").strip().lower() not in {
        "0", "false", "off", "no",
    }
    app.state.ingestion_worker = IngestionWorker(
        store_factory=get_store,
        poll_interval=float(os.getenv("INGESTION_WORKER_POLL_SECONDS", "1")),
        enabled=ingestion_worker_enabled,
        disabled_reason=("disabled_by_environment" if not ingestion_worker_enabled else ""),
    )
    memory_embedding_enabled = os.getenv("MEMORY_EMBEDDING", "off").strip().lower() in {
        "1", "true", "on", "yes",
    }
    app.state.memory_embedding_worker = MemoryEmbeddingWorker(
        store_factory=get_store,
        poll_interval=float(os.getenv("MEMORY_EMBEDDING_WORKER_POLL_SECONDS", "2")),
        batch_size=int(os.getenv("MEMORY_EMBEDDING_BATCH_SIZE", "20")),
        max_attempts=int(os.getenv("MEMORY_EMBEDDING_MAX_ATTEMPTS", "3")),
        retry_delay_seconds=int(os.getenv("MEMORY_EMBEDDING_RETRY_SECONDS", "30")),
        enabled=memory_embedding_enabled,
        disabled_reason=("MEMORY_EMBEDDING is off" if not memory_embedding_enabled else ""),
    )

    def scheduled_daily_journals(report_date: date | None = None) -> list[dict[str, Any]]:
        active_date = report_date or datetime.now(ZoneInfo("Asia/Shanghai")).date()
        results: list[dict[str, Any]] = []
        store = get_store()
        for context in list_project_contexts(store):
            results.extend(compile_project_journals(
                store=store,
                project_id=context.project_id,
                report_date=active_date,
                ingestion_wake=app.state.ingestion_worker.wake,
            ))
        return results

    def startup_daily_journal_backfill() -> None:
        previous_day = datetime.now(ZoneInfo("Asia/Shanghai")).date() - timedelta(days=1)
        try:
            scheduled_daily_journals(previous_day)
        except Exception as exc:
            app.state.journal_scheduler.record_error(f"startup backfill failed: {type(exc).__name__}: {exc}")

    def startup_daily_journal_recovery() -> None:
        try:
            get_store().recover_interrupted_daily_journals()
        except Exception as exc:
            app.state.journal_scheduler.record_error(f"startup recovery failed: {type(exc).__name__}: {exc}")

    app.state.journal_scheduler = DailyJournalScheduler(scheduled_daily_journals)
    inspection_worker_enabled = os.getenv("INSPECTION_WORKER_ENABLED", "on").strip().lower() not in {
        "0", "false", "off", "no",
    }
    app.state.inspection_worker = InspectionWorker(
        store_factory=get_store,
        runtime_factory=get_runtime,
        milestone_loader=lambda store, project_id: load_project_context(
            store,
            project_id,
        ).default_milestone,
        poll_interval=float(os.getenv("INSPECTION_WORKER_POLL_SECONDS", "1")),
        enabled=inspection_worker_enabled,
    )

    def scheduled_inspections() -> list[dict[str, Any]]:
        store = get_store()
        dispatcher = InspectionDispatcher(store)
        jobs = [
            dispatcher.enqueue_scheduled(project_id=context.project_id)
            for context in list_project_contexts(store)
        ]
        app.state.inspection_worker.wake()
        return jobs

    app.state.inspection_scheduler = InspectionScheduler(
        scheduled_inspections,
        hour=int(os.getenv("INSPECTION_SCHEDULE_HOUR", "23")),
        minute=int(os.getenv("INSPECTION_SCHEDULE_MINUTE", "40")),
    )
    app.state.conversation_runs = ConversationRunManager(
        store_factory=get_store,
        runtime_factory=get_runtime,
        milestone_loader=lambda store, project_id: load_project_context(
            store,
            project_id,
        ).default_milestone,
        ingestion_wake=app.state.ingestion_worker.wake,
        inspection_wake=app.state.inspection_worker.wake,
        final_transform=attach_scheduler_status,
        evolution_service_factory=MemorySkillEvolutionService,
        max_workers=int(os.getenv("CONVERSATION_RUN_WORKERS", "4")),
    )
    app.include_router(create_conversation_router(ConversationRouterServices(
        store_factory=get_store,
        run_manager=app.state.conversation_runs,
    )))
    app.include_router(create_skills_router(SkillRouterServices(
        store_factory=get_store,
        skill_service_factory=get_skill_service,
    )))
    app.include_router(create_memory_skill_router(MemorySkillRouterServices(
        store_factory=get_store,
        service_factory=MemorySkillEvolutionService,
    )))
    app.include_router(create_project_files_router(ProjectFileRouterServices(
        store_factory=get_store,
        ingestion_wake=app.state.ingestion_worker.wake,
    )))
    startup_hooks.extend([
        app.state.brief_scheduler.start,
        app.state.ingestion_worker.start,
        app.state.memory_embedding_worker.start,
        app.state.inspection_worker.start,
        app.state.inspection_scheduler.start,
        app.state.conversation_runs.recover,
        startup_daily_journal_recovery,
        app.state.journal_scheduler.start,
        startup_daily_journal_backfill,
    ])
    shutdown_hooks.extend([
        app.state.journal_scheduler.shutdown,
        app.state.conversation_runs.shutdown,
        app.state.inspection_scheduler.shutdown,
        app.state.inspection_worker.shutdown,
        app.state.memory_embedding_worker.shutdown,
        app.state.ingestion_worker.shutdown,
        app.state.brief_scheduler.shutdown,
    ])

    @app.get("/api/health")
    def health() -> dict[str, Any]:
        scheduler_status = app.state.brief_scheduler.status()
        journal_scheduler_status = app.state.journal_scheduler.status()
        return {
            "ok": True,
            "store_dir": str(root_dir),
            "runtime": "model-function-calling-harness",
            "frontend": "react-vite",
            "scheduler": scheduler_status,
            "daily_brief_scheduler": scheduler_status,
            "daily_journal_scheduler": journal_scheduler_status,
            "ingestion_worker": app.state.ingestion_worker.status(),
            "memory_embedding_worker": app.state.memory_embedding_worker.status(),
            "conversation_runtime": app.state.conversation_runs.status(),
            "inspection_worker": app.state.inspection_worker.status(),
            "inspection_scheduler": app.state.inspection_scheduler.status(),
        }

    @app.get("/api/ingestion/jobs")
    def ingestion_jobs(request: Request) -> dict[str, Any]:
        scope = request_scope(request)
        jobs = get_store().list_ingestion_jobs(
            project_id=scope.project_id,
            actor=scope.actor,
            ctx=scope.access_context,
            limit=200,
        )
        return {
            "jobs": [ingestion_job_view(job) for job in jobs],
            "worker": app.state.ingestion_worker.status(),
        }

    @app.get("/api/inspections")
    def inspection_jobs(request: Request) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        return {
            "jobs": get_store().inspection_jobs.list(project_id=scope.project_id),
            "worker": app.state.inspection_worker.status(),
            "scheduler": app.state.inspection_scheduler.status(),
        }

    @app.post("/api/inspections", status_code=202)
    def request_inspection(
        request: Request,
        payload: InspectionRequest,
    ) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        store = get_store()
        for source_id in payload.source_ids:
            source = store.get_source(source_id)
            if source is None:
                raise HTTPException(status_code=404, detail=f"Unknown source: {source_id}")
            require_visible(scope, source)
        job = InspectionDispatcher(store).enqueue_manual(
            actor=scope.actor,
            project_id=scope.project_id,
            source_ids=payload.source_ids,
            reason=payload.reason,
            request_id=payload.request_id,
        )
        app.state.inspection_worker.wake()
        return {"job": job, "worker": app.state.inspection_worker.status()}

    @app.post("/api/ingestion/jobs/{job_id}/retry")
    def retry_ingestion_job(request: Request, job_id: str) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        store = get_store()
        try:
            target = store.get_ingestion_job(job_id)
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown ingestion job: {job_id}")
        if target["project_id"] != scope.project_id:
            raise HTTPException(status_code=404, detail=f"Unknown ingestion job: {job_id}")
        require_visible(scope, target)
        try:
            job = store.retry_ingestion_job(job_id)
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        app.state.ingestion_worker.wake()
        return {
            "job": ingestion_job_view(job),
            "worker": app.state.ingestion_worker.status(),
        }

    @app.get("/api/identity/actor")
    def identity_actor(request: Request) -> dict[str, Any]:
        try:
            actor = resolve_actor(request, user_lookup=get_store().get_access_user)
        except IdentityResolutionError as exc:
            raise HTTPException(status_code=401, detail=str(exc)) from exc
        return {
            "id": actor.id,
            "org_id": actor.org_id,
            "name": actor.name,
            "feishu_id": actor.feishu_id,
        }

    @app.get("/api/dev/switchable-users")
    def switchable_users(request: Request) -> dict[str, Any]:
        scope = request_scope(request)
        return {
            "project_id": scope.project_id,
            "users": get_store().list_project_access_users(scope.project_id),
        }

    @app.get("/api/workspace")
    def workspace(request: Request) -> dict[str, Any]:
        return latest_workspace(request)

    @app.post("/api/daily-journal/generate")
    def daily_journal_generate(request: Request, payload: DailyJournalGenerateRequest) -> dict[str, Any]:
        scope = request_scope(request)
        if payload.report_date:
            try:
                report_date = date.fromisoformat(payload.report_date)
            except ValueError as exc:
                raise HTTPException(status_code=400, detail="report_date must use YYYY-MM-DD") from exc
        else:
            report_date = datetime.now(ZoneInfo("Asia/Shanghai")).date()
        try:
            journal = DailyJournalService(
                get_store(),
                ingestion_wake=app.state.ingestion_worker.wake,
            ).compile_day(
                project_id=scope.project_id,
                actor=scope.actor,
                report_date=report_date,
            )
        except DailyJournalValidationError as exc:
            raise HTTPException(status_code=502, detail=f"日报整理失败：{exc}") from exc
        return {"journal": journal, "workspace": latest_workspace(request)}

    @app.get("/api/daily-work-records")
    def daily_work_records(
        request: Request,
        date_from: str = "",
        date_to: str = "",
        kind: str = "",
        subject_user_id: str = "",
    ) -> dict[str, Any]:
        scope = request_scope(request)
        rows = get_store().list_daily_work_records(
            project_id=scope.project_id,
            actor=scope.actor,
            access_context=scope.access_context,
            date_from=date_from,
            date_to=date_to,
            kind=kind,
            subject_user_id=subject_user_id,
        )
        return {"records": rows, "count": len(rows)}

    @app.patch("/api/daily-work-records/{record_id}")
    def edit_daily_work_record(
        request: Request,
        record_id: str,
        payload: DailyWorkRecordEditRequest,
    ) -> dict[str, Any]:
        scope = request_scope(request)
        try:
            record = get_store().update_daily_work_record(
                record_id,
                actor=scope.actor,
                access_context=scope.access_context,
                patch=payload.model_dump(exclude_none=True),
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown daily work record: {record_id}") from exc
        except PermissionError as exc:
            raise HTTPException(status_code=403, detail=str(exc)) from exc
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        return {"record": record}

    @app.get("/api/project-skills")
    def project_skills_list(request: Request, published_only: bool = False) -> dict[str, Any]:
        scope = request_scope(request)
        skills = get_skill_service(get_store()).list_skills(
            scope.actor,
            scope.access_context,
            project_id=scope.project_id,
            published_only=published_only,
        )
        return {"skills": skills}

    @app.post("/api/project-skills/compile")
    def project_skill_compile(request: Request, payload: ProjectSkillCompileRequest) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        try:
            skill = get_skill_service(get_store()).compile_method(
                payload.method_id,
                scope.actor,
                scope.access_context,
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown visible method: {payload.method_id}") from exc
        except ProjectSkillValidationError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        return {"skill": skill, "workspace": latest_workspace(request)}

    @app.post("/api/project-skills/{skill_id}/test")
    def project_skill_test(request: Request, skill_id: str) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        try:
            result = get_skill_service(get_store()).run_pressure_tests(
                skill_id,
                scope.actor,
                scope.access_context,
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown visible Skill: {skill_id}") from exc
        except (ProjectSkillValidationError, ValueError) as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        return {"skill": result, "workspace": latest_workspace(request)}

    @app.post("/api/project-skills/{skill_id}/publish")
    def project_skill_publish(request: Request, skill_id: str) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        try:
            result = get_skill_service(get_store()).publish(
                skill_id,
                scope.actor,
                scope.access_context,
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown visible Skill: {skill_id}") from exc
        except (ProjectSkillValidationError, ValueError) as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        return {"skill": result, "workspace": latest_workspace(request)}

    @app.post("/api/workspace/archive-and-clear")
    def workspace_archive_and_clear(request: Request, payload: WorkspaceArchiveRequest) -> dict[str, Any]:
        store = get_store()
        scope = request_scope(request)
        require_role(scope, "pmo")
        actor = scope.actor
        project_id = scope.project_id
        if payload.confirmation != "ARCHIVE":
            raise HTTPException(status_code=400, detail="确认文本不正确，请输入 ARCHIVE。")
        archive = store.archive_project_workspace(project_id, actor_id=actor.id)
        preserved_sources = len([
            source
            for source in store.list_sources()
            if source.get("project_id") == project_id
        ])
        return {
            "archive": archive,
            "preserved_sources": preserved_sources,
            "workspace": latest_workspace(request),
        }

    @app.get("/api/brief/latest")
    def daily_brief_latest(request: Request) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        store = get_store()
        milestone = load_project_context(store, scope.project_id).default_milestone
        brief = latest_daily_brief(store, project_id=scope.project_id) or ensure_today_daily_brief(
            store,
            milestone,
            project_id=scope.project_id,
        )
        return {"brief": brief}

    @app.post("/api/brief/generate")
    def daily_brief_generate(request: Request, payload: BriefGenerateRequest) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        store = get_store()
        brief = generate_daily_brief(
            store=store,
            milestone=load_project_context(store, scope.project_id).default_milestone,
            project_id=scope.project_id,
            push_feishu=payload.push_feishu,
            webhook_url=os.getenv("FEISHU_WEBHOOK_URL", ""),
        )
        return {"brief": brief, "workspace": latest_workspace(request)}

    @app.post("/api/weekly-review/generate")
    def weekly_review_generate(request: Request, payload: WeeklyReviewGenerateRequest) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        week_end = _parse_date(payload.week_end) if payload.week_end else None
        if payload.week_end and week_end is None:
            raise HTTPException(status_code=400, detail="week_end 必须是 YYYY-MM-DD。")
        store = get_store()
        review = generate_weekly_review(
            store,
            week_end=week_end,
            project_id=scope.project_id,
            actor=scope.actor,
            access_context=scope.access_context,
        )
        return {"weekly_review": review, "workspace": latest_workspace(request)}

    @app.post("/api/meetings/continuations")
    def meeting_continuation_mark(request: Request, payload: MeetingContinuationRequest) -> dict[str, Any]:
        store = get_store()
        scope = request_scope(request)
        require_role(scope, "topic_lead")
        for source_id in (payload.previous_source_id, payload.next_source_id):
            source = store.get_source(source_id)
            if source is None:
                raise HTTPException(status_code=404, detail=f"Unknown source: {source_id}")
            require_visible(scope, source)
        try:
            result = mark_meeting_continuation(
                store,
                previous_source_id=payload.previous_source_id,
                next_source_id=payload.next_source_id,
                editor=scope.actor.id,
                notes=payload.notes,
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown source: {exc}") from exc
        return {"continuation": result}

    @app.get("/api/meetings/{source_id}/followup-comparison")
    def meeting_followup_comparison(request: Request, source_id: str) -> dict[str, Any]:
        store = get_store()
        scope = request_scope(request)
        source = store.get_source(source_id)
        if source is None:
            raise HTTPException(status_code=404, detail=f"Unknown source: {source_id}")
        require_visible(scope, source)
        try:
            comparison = build_meeting_followup_comparison(
                store,
                next_source_id=source_id,
                project_id=scope.project_id,
                actor=scope.actor,
                access_context=scope.access_context,
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown source: {exc}") from exc
        return {"comparison": comparison}

    @app.get("/api/file-workspace")
    def file_workspace(request: Request) -> dict[str, Any]:
        require_role(request_scope(request), "pmo")
        return get_file_workspace().status()

    @app.post("/api/file-workspace/steps/{step_id}/run")
    def run_file_workspace_step(request: Request, step_id: str) -> dict[str, Any]:
        require_role(request_scope(request), "pmo")
        workspace = get_file_workspace()
        try:
            workspace.run_step(step_id)
        except PermissionError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        except ValueError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
        return workspace.status()

    @app.post("/api/file-workspace/steps/{step_id}/confirm")
    def confirm_file_workspace_step(request: Request, step_id: str, payload: ConfirmationRequest) -> dict[str, Any]:
        require_role(request_scope(request), "pmo")
        workspace = get_file_workspace()
        try:
            return workspace.confirm_step(step_id, notes=payload.notes)
        except PermissionError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        except ValueError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc

    @app.post("/api/milestones/config")
    def update_milestone_config(request: Request, payload: MilestoneConfigRequest) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        update_project_context(
            get_store(),
            scope.project_id,
            payload.model_dump(exclude_none=True),
            actor_id=scope.actor.id,
        )
        return latest_workspace(request)

    @app.get("/api/deliverables")
    def list_deliverables(request: Request, milestone_id: str = "") -> dict[str, Any]:
        scope = request_scope(request)
        rows = get_deliverable_service().list_requirements(
            actor=scope.actor,
            access_context=scope.access_context,
            project_id=scope.project_id,
            milestone_id=milestone_id,
        )
        return {"deliverables": [_public_deliverable(row) for row in rows], "count": len(rows)}

    @app.post("/api/deliverables", status_code=201)
    def create_deliverable(
        request: Request,
        payload: DeliverableRequirementRequest,
    ) -> dict[str, Any]:
        scope = request_scope(request)
        milestone = load_project_context(get_store(), scope.project_id).default_milestone
        try:
            row = get_deliverable_service().create_requirement(
                actor=scope.actor,
                access_context=scope.access_context,
                project_id=scope.project_id,
                milestone_id=milestone.milestone_id,
                **payload.model_dump(),
            )
        except PermissionError as exc:
            raise HTTPException(status_code=403, detail=str(exc)) from exc
        except (DeliverableValidationError, ValueError) as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        return {"deliverable": _public_deliverable(row)}

    @app.patch("/api/deliverables/{deliverable_id}")
    def update_deliverable(
        request: Request,
        deliverable_id: str,
        payload: DeliverableRequirementEditRequest,
    ) -> dict[str, Any]:
        scope = request_scope(request)
        try:
            row = get_deliverable_service().update_requirement(
                deliverable_id,
                actor=scope.actor,
                access_context=scope.access_context,
                patch=payload.model_dump(exclude_none=True),
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown deliverable: {deliverable_id}") from exc
        except PermissionError as exc:
            raise HTTPException(status_code=403, detail=str(exc)) from exc
        except (DeliverableValidationError, ValueError) as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        return {"deliverable": _public_deliverable(row)}

    @app.post("/api/deliverables/{deliverable_id}/versions", status_code=201)
    async def submit_deliverable_version(
        request: Request,
        deliverable_id: str,
        files: list[UploadFile] = File(...),
        artifact_roles: str = Form(""),
        note: str = Form(""),
    ) -> dict[str, Any]:
        scope = request_scope(request)
        if not files:
            raise HTTPException(status_code=400, detail="至少需要一个交付物文件。")
        if len(files) > MAX_DELIVERABLE_FILES:
            raise HTTPException(
                status_code=400,
                detail=f"一次最多上传 {MAX_DELIVERABLE_FILES} 个文件。",
            )
        try:
            parsed_roles = json.loads(artifact_roles) if artifact_roles.strip() else []
        except json.JSONDecodeError as exc:
            raise HTTPException(status_code=400, detail="artifact_roles 必须是 JSON 数组。") from exc
        if not isinstance(parsed_roles, list):
            raise HTTPException(status_code=400, detail="artifact_roles 必须是 JSON 数组。")
        if parsed_roles and len(parsed_roles) != len(files):
            raise HTTPException(status_code=400, detail="artifact_roles 数量必须与文件数量一致。")
        try:
            file_size_limit = deliverable_file_size_limit()
            total_size_limit = deliverable_total_size_limit()
        except DeliverableValidationError as exc:
            raise HTTPException(status_code=500, detail=str(exc)) from exc
        uploads: list[DeliverableUpload] = []
        total_size = 0
        for index, upload in enumerate(files):
            filename = Path(upload.filename or "deliverable_file").name
            content = await _read_upload_with_limit(
                upload,
                file_size_limit,
                total_remaining=total_size_limit - total_size,
            )
            total_size += len(content)
            uploads.append(DeliverableUpload(
                filename=filename,
                content=content,
                content_type=upload.content_type or "",
                artifact_role=str(parsed_roles[index] if parsed_roles else "formal"),
            ))
        try:
            version = get_deliverable_service().submit_version(
                deliverable_id,
                actor=scope.actor,
                access_context=scope.access_context,
                files=uploads,
                note=note,
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown deliverable: {deliverable_id}") from exc
        except PermissionError as exc:
            raise HTTPException(status_code=403, detail=str(exc)) from exc
        except (DeliverableValidationError, ValueError) as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        return {"version": _public_deliverable_version(version)}

    @app.get("/api/deliverable-files/{file_id}/download")
    def download_deliverable_file(request: Request, file_id: str) -> FileResponse:
        scope = request_scope(request)
        store = get_store()
        try:
            metadata = store.get_deliverable_file(file_id)
            path = get_deliverable_service(store).resolve_download(
                file_id,
                actor=scope.actor,
                access_context=scope.access_context,
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown deliverable file: {file_id}") from exc
        except PermissionError as exc:
            raise HTTPException(status_code=403, detail=str(exc)) from exc
        except FileNotFoundError as exc:
            raise HTTPException(status_code=410, detail=str(exc)) from exc
        except DeliverableValidationError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        return FileResponse(
            path,
            filename=metadata["original_name"],
            media_type=metadata.get("content_type") or "application/octet-stream",
        )

    @app.post("/api/sources/upload", status_code=202)
    async def upload_source(
        request: Request,
        file: UploadFile | None = File(None),
        files: list[UploadFile] | None = File(None),
        raw_file: UploadFile | None = File(None),
        title: str = Form(""),
        meeting_date: str = Form(""),
        topic: str = Form(""),
        topic_id: str = Form(""),
        sensitivity: str = Form("l1"),
        input_kind: str = Form("auto"),
    ) -> dict[str, Any]:
        store = get_store()
        scope = request_scope(request)
        actor = scope.actor
        project_id = scope.project_id
        if topic_id and actor.id not in scope.access_context.members_of(topic_id):
            raise HTTPException(status_code=403, detail="Actor is not a member of the selected topic")
        ingest_tags = {
            "org_id": actor.org_id,
            "project_id": project_id,
            "topic_id": topic_id or None,
            "author_id": actor.id,
            "sensitivity": sensitivity,
        }
        uploaded_files = [*([file] if file is not None else []), *(files or [])]
        if not uploaded_files:
            raise HTTPException(status_code=400, detail="至少需要上传一个文件。")
        if len(uploaded_files) > 20:
            raise HTTPException(status_code=400, detail="一次最多上传 20 个文件。")
        if raw_file is not None and len(uploaded_files) != 1:
            raise HTTPException(status_code=400, detail="原始转写配对只支持单文件上传。")
        contents = [await upload.read() for upload in uploaded_files]
        if any(not content for content in contents):
            raise HTTPException(status_code=400, detail="上传文件不能为空。")
        raw_content = await raw_file.read() if raw_file else None
        try:
            prepared = [
                (
                    upload,
                    content,
                    resolve_project_input_kind(upload.filename or "project_file", content, input_kind),
                )
                for upload, content in zip(uploaded_files, contents)
            ]
            queued: list[tuple[Any, dict[str, Any]]] = []
            for upload, content, resolved_kind in prepared:
                filename = upload.filename or "project_file"
                source = register_uploaded_project_file(
                    filename=filename,
                    content=content,
                    title=title if len(prepared) == 1 else Path(filename).stem,
                    meeting_date=meeting_date,
                    topic=topic,
                    access_tags=ingest_tags,
                    raw_filename=raw_file.filename if raw_file else "",
                    raw_content=raw_content,
                    input_kind=resolved_kind,
                )
                store.ingest(
                    source,
                    tags=ingest_tags,
                    actor=actor,
                    kind=source.input_kind,
                    materialize=False,
                )
                job = enqueue_source_ingestion(store=store, source=source, actor=actor)
                queued.append((source, job))
        except (ValueError, MeetingIngestionError) as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        app.state.ingestion_worker.wake()
        source, job = queued[0]
        return {
            "source": source_ingestion_view(source),
            "job": ingestion_job_view(job),
            "sources": [source_ingestion_view(item) for item, _ in queued],
            "jobs": [ingestion_job_view(item) for _, item in queued],
            "worker": app.state.ingestion_worker.status(),
        }

    @app.get("/api/assets/people")
    def people_asset_status(request: Request) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        structure = load_project_people(
            get_store(),
            scope.project_id,
            actor=scope.actor,
            access_context=scope.access_context,
        )
        source_path = project_people_asset_path(
            get_store(),
            scope.project_id,
            actor=scope.actor,
            access_context=scope.access_context,
        )
        return _people_asset_payload(
            structure,
            source_path or _project_people_asset_path(root_dir, scope.project_id),
        )

    @app.post("/api/assets/people/upload")
    async def upload_people_asset(request: Request, file: UploadFile = File(...)) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        content = await file.read()
        if not content:
            raise HTTPException(status_code=400, detail="上传文件为空。")
        suffix = Path(file.filename or "").suffix.lower()
        if suffix not in {".csv", ".xlsx", ".xlsm", ".xmind"}:
            raise HTTPException(
                status_code=400,
                detail="人员资产只支持 .csv、.xlsx、.xlsm 或 .xmind。",
            )
        upload_dir = root_dir / "project_assets" / scope.project_id / "uploads"
        upload_dir.mkdir(parents=True, exist_ok=True)
        original_stem = Path(file.filename or "people_asset").stem
        safe_stem = "".join(
            character if character.isalnum() or character in "._-" else "_"
            for character in original_stem
        ).strip("._")[:80] or "people_asset"
        target = upload_dir / f"{sha256(content).hexdigest()}-{safe_stem}{suffix}"
        temporary = upload_dir / f".{uuid4().hex}.upload{suffix}"
        temporary.write_bytes(content)
        try:
            if suffix == ".xmind":
                structure = load_people_structure_from_xmind(temporary)
            else:
                structure = load_people_structure_from_table(temporary)
            if target.exists():
                temporary.unlink()
            else:
                os.replace(temporary, target)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        finally:
            temporary.unlink(missing_ok=True)
        asset_path = _persist_project_people(
            get_store(),
            scope,
            structure,
            root_dir=root_dir,
            source_path=target,
            source_type=f"{suffix.removeprefix('.')} import",
        )
        return {
            "people_asset": _people_asset_payload(structure, asset_path),
            "workspace": latest_workspace(request),
        }

    @app.get("/api/assets/people/template")
    def download_people_import_template(request: Request) -> StreamingResponse:
        scope = request_scope(request)
        require_role(scope, "pmo")
        content = _build_people_import_template()
        return StreamingResponse(
            BytesIO(content),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="people_import_template.xlsx"'},
        )

    @app.post("/api/assets/people/import")
    async def import_people_from_template(request: Request, file: UploadFile = File(...)) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        content = await file.read()
        if not content:
            raise HTTPException(status_code=400, detail="导入文件为空。")
        suffix = Path(file.filename or "").suffix.lower()
        if suffix not in {".csv", ".xlsx", ".xlsm"}:
            raise HTTPException(status_code=400, detail="批量导入只支持模板 Excel 或 CSV 文件。")
        upload_dir = root_dir / "project_assets" / scope.project_id / "imports"
        upload_dir.mkdir(parents=True, exist_ok=True)
        safe_name = "".join(
            character if character.isalnum() or character in "._-" else "_"
            for character in Path(file.filename or f"people_import{suffix}").name
        )[:100]
        target = upload_dir / f"{sha256(content).hexdigest()}-{safe_name}"
        temporary = upload_dir / f".{uuid4().hex}.import{suffix}"
        temporary.write_bytes(content)
        try:
            imported = load_people_structure_from_table(temporary)
            if not imported.people and not imported.groups:
                raise ValueError("模板中没有可导入的人员或组别，请填写后再导入。")
            current = load_project_people(
                get_store(),
                scope.project_id,
                actor=scope.actor,
                access_context=scope.access_context,
            )
            merged, summary = _merge_people_structures(current, imported)
            if target.exists():
                temporary.unlink()
            else:
                os.replace(temporary, target)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        finally:
            temporary.unlink(missing_ok=True)
        asset_path = _persist_project_people(
            get_store(),
            scope,
            merged,
            root_dir=root_dir,
            source_path=target,
            source_type=f"{suffix.removeprefix('.')} template import",
        )
        return {
            "people_asset": _people_asset_payload(merged, asset_path),
            "import_summary": summary,
            "workspace": latest_workspace(request),
        }

    @app.post("/api/items/{item_id}/confirm")
    def confirm_item(request: Request, item_id: str, payload: ConfirmationRequest) -> dict[str, Any]:
        store = get_store()
        try:
            item = store.get_item(item_id)
            require_review(request_scope(request), item)
            item = store.confirm_item(item_id, editor=request_scope(request).actor.id, notes=payload.notes)
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown item: {item_id}") from exc
        return to_plain(item)

    @app.post("/api/items/{item_id}/reject")
    def reject_item(request: Request, item_id: str, payload: ConfirmationRequest) -> dict[str, Any]:
        store = get_store()
        try:
            item = store.get_item(item_id)
            require_review(request_scope(request), item)
            item = store.reject_item(item_id, editor=request_scope(request).actor.id, notes=payload.notes)
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown item: {item_id}") from exc
        return to_plain(item)

    @app.post("/api/items/{item_id}/archive-task-candidate")
    def archive_task_candidate(request: Request, item_id: str, payload: ConfirmationRequest) -> dict[str, Any]:
        store = get_store()
        try:
            item = store.get_item(item_id)
            require_review(request_scope(request), item)
            if item.category not in {"task", "tasks", "followup"}:
                raise HTTPException(status_code=400, detail="Only task candidates can use this endpoint")
            if item.status != CandidateStatus.CANDIDATE:
                raise HTTPException(status_code=400, detail="Only pending task candidates can be deleted")
            store.update_item_fields(
                item_id,
                status=CandidateStatus.ARCHIVED.value,
                editor=request_scope(request).actor.id,
                notes=payload.notes or "人工判断无需形成正式任务",
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown item: {item_id}") from exc
        return latest_workspace(request)

    @app.post("/api/items/{item_id}/publish-task")
    def publish_item_as_task(request: Request, item_id: str, payload: ConfirmationRequest) -> dict[str, Any]:
        store = get_store()
        latest = store.latest_run(project_id=request_scope(request).project_id)
        try:
            require_review(request_scope(request), store.get_item(item_id))
            store.publish_item_as_work_item(
                item_id,
                editor=request_scope(request).actor.id,
                notes=payload.notes,
                milestone_id=latest.milestone_id if latest else "",
                source_run_id=latest.run_id if latest else "",
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown item: {item_id}") from exc
        return latest_workspace(request)

    @app.post("/api/work-items/{work_item_id}")
    def update_work_item(request: Request, work_item_id: str, payload: WorkItemEditRequest) -> dict[str, Any]:
        store = get_store()
        try:
            stored_task = store.get_work_item(work_item_id)
            require_task_edit(request_scope(request), stored_task)
            if payload.status is not None and payload.status != stored_task.status.value:
                require_task_status_edit(request_scope(request), stored_task, payload.status)
            supplied_fields = (
                payload.model_fields_set
                if hasattr(payload, "model_fields_set")
                else getattr(payload, "__fields_set__", set())
            )
            progress_update = (
                {"progress_percent": payload.progress_percent}
                if "progress_percent" in supplied_fields
                else {}
            )
            store.update_work_item_fields(
                work_item_id,
                title=payload.title,
                description=payload.description,
                status=payload.status,
                due_date=payload.due_date,
                deliverable=payload.deliverable,
                acceptance_criteria=payload.acceptance_criteria,
                professional_id=payload.professional_id,
                board_id=payload.board_id,
                planned_start=payload.planned_start,
                owner_candidates=payload.owner_candidates,
                linked_issue_ids=payload.linked_issue_ids,
                linked_task_ids=payload.linked_task_ids,
                editor=request_scope(request).actor.id,
                notes=payload.notes,
                **progress_update,
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown work item: {work_item_id}") from exc
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        return latest_workspace(request)

    @app.post("/api/three-lists/issues")
    def create_three_list_issue(request: Request, payload: ThreeListMutationRequest) -> dict[str, Any]:
        store = get_store()
        scope = request_scope(request)
        require_role(scope, "topic_lead")
        actor = scope.actor
        _create_manual_item(store, payload, category="issue", prefix="manual_issue", actor=actor)
        return latest_workspace(request)

    @app.post("/api/three-lists/issues/{item_id}")
    def update_three_list_issue(request: Request, item_id: str, payload: ThreeListMutationRequest) -> dict[str, Any]:
        store = get_store()
        try:
            require_item_edit(request_scope(request), store.get_item(item_id))
            _update_item_from_three_list(store, item_id, payload, editor=request_scope(request).actor.id)
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown issue: {item_id}") from exc
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        return latest_workspace(request)

    @app.post("/api/three-lists/issues/{item_id}/archive")
    def archive_three_list_issue(request: Request, item_id: str, payload: ConfirmationRequest | None = None) -> dict[str, Any]:
        store = get_store()
        try:
            require_item_edit(request_scope(request), store.get_item(item_id))
            store.update_item_fields(
                item_id,
                status=CandidateStatus.ARCHIVED.value,
                editor=request_scope(request).actor.id,
                notes=(payload.notes if payload else "从三清单归档"),
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown issue: {item_id}") from exc
        return latest_workspace(request)

    @app.post("/api/three-lists/tasks")
    def create_three_list_task(request: Request, payload: ThreeListMutationRequest) -> dict[str, Any]:
        store = get_store()
        scope = request_scope(request)
        require_role(scope, "topic_lead")
        actor = scope.actor
        item = _create_manual_item(store, payload, category="followup", prefix="manual_task", actor=actor)
        latest = store.latest_run(project_id=request_scope(request).project_id)
        work_item = store.publish_item_as_work_item(
            item.item_id,
            editor=actor.id,
            notes=payload.notes or "三清单人工新增任务",
            milestone_id=(
                latest.milestone_id
                if latest
                else load_project_context(store, request_scope(request).project_id).default_milestone.milestone_id
            ),
            source_run_id=latest.run_id if latest else "",
        )
        store.update_work_item_fields(
            work_item.work_item_id,
            title=payload.title,
            description=payload.description,
            status=payload.status or "open",
            due_date=payload.due_date,
            deliverable=payload.deliverable,
            acceptance_criteria=payload.acceptance_criteria,
            owner_candidates=payload.owner_candidates,
            linked_issue_ids=payload.linked_issue_ids,
            linked_task_ids=payload.linked_task_ids,
            editor=actor.id,
            notes=payload.notes or "三清单人工新增任务",
        )
        return latest_workspace(request)

    @app.post("/api/three-lists/tasks/{task_id}")
    def update_three_list_task(request: Request, task_id: str, payload: ThreeListMutationRequest) -> dict[str, Any]:
        store = get_store()
        try:
            stored_task = store.get_work_item(task_id)
            require_task_edit(request_scope(request), stored_task)
            if payload.status is not None and payload.status != stored_task.status.value:
                require_task_status_edit(request_scope(request), stored_task, payload.status)
            store.update_work_item_fields(
                task_id,
                title=payload.title,
                description=payload.description,
                status=payload.status,
                due_date=payload.due_date,
                deliverable=payload.deliverable,
                acceptance_criteria=payload.acceptance_criteria,
                owner_candidates=payload.owner_candidates,
                linked_issue_ids=payload.linked_issue_ids,
                linked_task_ids=payload.linked_task_ids,
                editor=request_scope(request).actor.id,
                notes=payload.notes,
            )
        except KeyError:
            try:
                candidate = store.get_item(task_id)
                scope = request_scope(request)
                if payload.status is not None and payload.status != candidate.status.value:
                    require_task_status_edit(scope, candidate, payload.status)
                    latest = store.latest_run(project_id=scope.project_id)
                    work_item = store.publish_item_as_work_item(
                        task_id,
                        editor=scope.actor.id,
                        notes=payload.notes or "负责人确认候选任务并更新状态",
                        milestone_id=latest.milestone_id if latest else "",
                        source_run_id=latest.run_id if latest else "",
                    )
                    store.update_work_item_fields(
                        work_item.work_item_id,
                        status=payload.status,
                        editor=scope.actor.id,
                        notes=payload.notes or "负责人更新任务状态",
                    )
                else:
                    require_item_edit(scope, candidate)
                    _update_item_from_three_list(store, task_id, payload, editor=scope.actor.id)
            except KeyError as exc:
                raise HTTPException(status_code=404, detail=f"Unknown task: {task_id}") from exc
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        return latest_workspace(request)

    @app.post("/api/three-lists/tasks/{task_id}/archive")
    def archive_three_list_task(request: Request, task_id: str, payload: ConfirmationRequest | None = None) -> dict[str, Any]:
        store = get_store()
        try:
            require_task_edit(request_scope(request), store.get_work_item(task_id))
            store.update_work_item_fields(
                task_id,
                status="archived",
                editor=request_scope(request).actor.id,
                notes=(payload.notes if payload else "从三清单归档"),
            )
        except KeyError:
            try:
                require_item_edit(request_scope(request), store.get_item(task_id))
                store.update_item_fields(
                    task_id,
                    status=CandidateStatus.ARCHIVED.value,
                    editor=request_scope(request).actor.id,
                    notes=(payload.notes if payload else "从三清单归档"),
                )
            except KeyError as exc:
                raise HTTPException(status_code=404, detail=f"Unknown task: {task_id}") from exc
        return latest_workspace(request)

    @app.post("/api/three-lists/methods")
    def create_three_list_method(request: Request, payload: ThreeListMutationRequest) -> dict[str, Any]:
        store = get_store()
        scope = request_scope(request)
        require_role(scope, "topic_lead")
        actor = scope.actor
        _create_manual_item(store, payload, category="methods", prefix="manual_method", actor=actor)
        return latest_workspace(request)

    @app.post("/api/three-lists/methods/{item_id}")
    def update_three_list_method(request: Request, item_id: str, payload: ThreeListMutationRequest) -> dict[str, Any]:
        store = get_store()
        try:
            require_item_edit(request_scope(request), store.get_item(item_id))
            _update_item_from_three_list(store, item_id, payload, editor=request_scope(request).actor.id)
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown method: {item_id}") from exc
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        return latest_workspace(request)

    @app.post("/api/three-lists/methods/{item_id}/archive")
    def archive_three_list_method(request: Request, item_id: str, payload: ConfirmationRequest | None = None) -> dict[str, Any]:
        store = get_store()
        try:
            require_item_edit(request_scope(request), store.get_item(item_id))
            store.update_item_fields(
                item_id,
                status=CandidateStatus.ARCHIVED.value,
                editor=request_scope(request).actor.id,
                notes=(payload.notes if payload else "从三清单归档"),
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown method: {item_id}") from exc
        return latest_workspace(request)

    @app.post("/api/people/groups")
    def create_people_group(request: Request, payload: PeopleGroupRequest) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        structure = load_project_people(
            get_store(), scope.project_id, actor=scope.actor, access_context=scope.access_context
        )
        name = payload.name.strip()
        if not name:
            raise HTTPException(status_code=400, detail="组别名称不能为空。")
        if any(group.name == name for group in structure.groups):
            raise HTTPException(status_code=409, detail=f"组别已存在：{name}")
        structure.groups.append(PeopleGroup(name=name, path=(payload.path or _person_path(structure.root_title, name, "", "")).strip()))
        _persist_project_people(
            get_store(), scope, structure, root_dir=root_dir, source_path=None, source_type="manual group create"
        )
        return latest_workspace(request)

    @app.post("/api/people/groups/{group_ref}")
    def update_people_group(request: Request, group_ref: str, payload: PeopleGroupRequest) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        structure = load_project_people(
            get_store(), scope.project_id, actor=scope.actor, access_context=scope.access_context
        )
        target = next((group for group in structure.groups if group.name == group_ref), None)
        if target is None:
            raise HTTPException(status_code=404, detail=f"未找到组别：{group_ref}")
        name = payload.name.strip()
        if not name:
            raise HTTPException(status_code=400, detail="组别名称不能为空。")
        if name != group_ref and any(group.name == name for group in structure.groups):
            raise HTTPException(status_code=409, detail=f"组别已存在：{name}")
        old_name = target.name
        target.name = name
        target.path = (payload.path or _person_path(structure.root_title, name, "", "")).strip()
        for person in structure.people:
            if person.group == old_name:
                person.group = name
                person.path = _person_path(structure.root_title, name, person.role, person.name)
            for assignment in person.assignments:
                if assignment.group == old_name:
                    assignment.group = name
                    assignment.path = _person_path(structure.root_title, name, assignment.role, person.name)
        for scenario in structure.scenarios:
            scenario.teams = [name if team == old_name else team for team in scenario.teams]
        _persist_project_people(
            get_store(), scope, structure, root_dir=root_dir, source_path=None, source_type="manual group update"
        )
        return latest_workspace(request)

    @app.post("/api/people/groups/{group_ref}/delete")
    def delete_people_group(
        request: Request,
        group_ref: str,
        payload: ConfirmationRequest | None = None,
    ) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        structure = load_project_people(
            get_store(), scope.project_id, actor=scope.actor, access_context=scope.access_context
        )
        if not any(group.name == group_ref for group in structure.groups):
            raise HTTPException(status_code=404, detail=f"未找到组别：{group_ref}")
        member_count = sum(1 for person in structure.people if person.group == group_ref)
        if member_count:
            raise HTTPException(
                status_code=409,
                detail=f"组别“{group_ref}”仍有{member_count}名人员，请先把人员调整到其他板块后再删除。",
            )
        structure.groups = [group for group in structure.groups if group.name != group_ref]
        for scenario in structure.scenarios:
            scenario.teams = [team for team in scenario.teams if team != group_ref]
        _persist_project_people(
            get_store(), scope, structure, root_dir=root_dir, source_path=None, source_type="manual group delete"
        )
        return latest_workspace(request)

    @app.post("/api/people/{person_ref}")
    def update_person_profile(request: Request, person_ref: str, payload: PersonProfileRequest) -> dict[str, Any]:
        scope = request_scope(request)
        require_role(scope, "pmo")
        store = get_store()
        structure = load_project_people(
            store,
            scope.project_id,
            actor=scope.actor,
            access_context=scope.access_context,
        )
        target = next((person for person in structure.people if person.person_id == person_ref), None)
        if target is None:
            name_matches = [person for person in structure.people if person.name == person_ref]
            if len(name_matches) > 1:
                raise HTTPException(status_code=409, detail="存在同名人员，请使用 person_id 选择具体人员。")
            target = name_matches[0] if name_matches else None
        if target is None:
            new_name = (payload.name or person_ref).strip()
            if not new_name:
                raise HTTPException(status_code=400, detail="Person name is required")
            target = Person(
                name=new_name,
                group=payload.group or "",
                role=payload.role or "",
                path=payload.path or _person_path(
                    structure.root_title,
                    payload.group or "",
                    payload.role or "",
                    new_name,
                ),
            )
            structure.people.append(target)
        elif payload.name is not None and payload.name.strip() != target.name:
            new_name = payload.name.strip()
            if not new_name:
                raise HTTPException(status_code=400, detail="Person name is required")
            if any(person is not target and person.name == new_name for person in structure.people):
                raise HTTPException(status_code=400, detail=f"Person already exists: {new_name}")
            target.name = new_name
        if payload.group is not None:
            target.group = payload.group
        if payload.role is not None:
            target.role = payload.role
        if payload.path is not None:
            target.path = payload.path
        if payload.responsibility_note is not None:
            target.responsibility_note = payload.responsibility_note
        if payload.path is None and (
            payload.name is not None or payload.group is not None or payload.role is not None
        ):
            target.path = _person_path(structure.root_title, target.group, target.role, target.name)
        if not target.path:
            target.path = _person_path(structure.root_title, target.group, target.role, target.name)
        if target.assignments:
            primary_assignment = target.assignments[0]
            primary_assignment.group = target.group
            primary_assignment.role = target.role
            primary_assignment.path = target.path
            primary_assignment.responsibility_note = target.responsibility_note
        if target.group and all(group.name != target.group for group in structure.groups):
            structure.groups.append(
                PeopleGroup(
                    name=target.group,
                    path=_person_path(structure.root_title, target.group, "", ""),
                )
            )
        if os.getenv("PROJECT_AGENT_LEGACY_FIXTURE_MODE", "").strip().lower() == "on":
            save_people_asset(
                structure,
                configured_people_asset_path(),
                source_type="manual_update",
                notes=f"Updated person profile by {scope.actor.id}. {payload.notes}".strip(),
            )
        else:
            _persist_project_people(
                store,
                scope,
                structure,
                root_dir=root_dir,
                source_path=None,
                source_type="manual_update",
            )
        return latest_workspace(request)

    @app.post("/api/items/{item_id}/edit")
    def edit_item(request: Request, item_id: str, payload: EditItemRequest) -> dict[str, Any]:
        store = get_store()
        try:
            require_item_edit(request_scope(request), store.get_item(item_id))
            store.update_item_fields(
                item_id,
                status=payload.status,
                title=payload.title,
                description=payload.description,
                due_date=payload.due_date,
                deliverable=payload.deliverable,
                acceptance_criteria=payload.acceptance_criteria,
                owner_candidates=payload.owner_candidates,
                linked_issue_ids=payload.linked_issue_ids,
                linked_task_ids=payload.linked_task_ids,
                editor=request_scope(request).actor.id,
                notes=payload.notes,
            )
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=f"Unknown item: {item_id}") from exc
        return latest_workspace(request)

    return app


def _build_people_import_template() -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:
        raise RuntimeError("生成导入模板需要 openpyxl。") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员导入"
    headers = ["类型", "姓名", "所属业务板块", "职责", "职责备注", "人员ID", "路径", "资产名称"]
    sheet.append(headers)
    for _ in range(50):
        sheet.append(["person", "", "", "", "", "", "", ""])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:H51"
    widths = [14, 16, 20, 20, 36, 24, 42, 24]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    record_type = DataValidation(type="list", formula1='"person,group"', allow_blank=False)
    sheet.add_data_validation(record_type)
    record_type.add("A2:A51")

    instructions = workbook.create_sheet("填写说明")
    instruction_rows = [
        ["字段", "填写要求"],
        ["类型", "人员填写 person；仅新增空组时填写 group。"],
        ["姓名", "person 行填写人员姓名；group 行填写组别名称。必填。"],
        ["所属业务板块", "person 行必填，必须与组织看板组别名称一致；不存在时会自动创建。"],
        ["职责", "人员角色或岗位职责，例如：专题负责人、产品设计、实施人员。"],
        ["职责备注", "具体负责范围、边界或补充说明。"],
        ["人员ID", "选填。更新已有人员时建议填写；留空则按姓名+板块+职责匹配。"],
        ["路径", "选填。留空时系统按资产名称/板块/职责/姓名自动生成。"],
        ["资产名称", "选填。仅在创建第一份人员资产时使用。"],
        ["导入方式", "增量合并：不会删除模板中未出现的已有人员。完全替换请使用“整表更新”。"],
    ]
    for row in instruction_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 18
    instructions.column_dimensions["B"].width = 90
    for cell in instructions[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _merge_people_structures(current: Any, imported: Any) -> tuple[Any, dict[str, int]]:
    group_by_name = {group.name: group for group in current.groups if group.name}
    groups_added = 0
    for group in imported.groups:
        if not group.name or group.name in group_by_name:
            continue
        current.groups.append(group)
        group_by_name[group.name] = group
        groups_added += 1

    by_id = {person.person_id: person for person in current.people if person.person_id}
    by_key = {
        (person.name.strip(), person.group.strip(), person.role.strip()): person
        for person in current.people
    }
    people_added = 0
    people_updated = 0
    for person in imported.people:
        key = (person.name.strip(), person.group.strip(), person.role.strip())
        target = by_id.get(person.person_id) if person.person_id else None
        target = target or by_key.get(key)
        if target is None:
            current.people.append(person)
            if person.person_id:
                by_id[person.person_id] = person
            by_key[key] = person
            people_added += 1
        else:
            target.name = person.name or target.name
            target.group = person.group or target.group
            target.role = person.role or target.role
            target.responsibility_note = person.responsibility_note
            target.path = person.path or _person_path(
                current.root_title, target.group, target.role, target.name
            )
            if person.assignments:
                target.assignments = person.assignments
            people_updated += 1
        if person.group and person.group not in group_by_name:
            group = PeopleGroup(
                name=person.group,
                path=_person_path(current.root_title, person.group, "", ""),
            )
            current.groups.append(group)
            group_by_name[group.name] = group
            groups_added += 1

    return current, {
        "people_added": people_added,
        "people_updated": people_updated,
        "groups_added": groups_added,
        "people_total": len(current.people),
        "groups_total": len(current.groups),
    }


def _people_asset_payload(structure: Any, path: Path) -> dict[str, Any]:
    return {
        "path": str(path),
        "root_title": structure.root_title,
        "group_count": len(structure.groups),
        "people_count": len(structure.people),
        "scenario_count": len(structure.scenarios),
        "scenario_ids": [scenario.scenario_id for scenario in structure.scenarios],
    }


def _public_deliverable(row: dict[str, Any]) -> dict[str, Any]:
    return {
        key: value
        for key, value in row.items()
        if key not in {"payload", "relative_path"}
        and key != "versions"
    } | {
        "versions": [_public_deliverable_version(version) for version in row.get("versions", [])]
    }


def _public_deliverable_version(row: dict[str, Any]) -> dict[str, Any]:
    return {
        key: value
        for key, value in row.items()
        if key not in {"payload", "relative_path", "files"}
    } | {
        "files": [
            {
                key: value
                for key, value in file_row.items()
                if key not in {"payload", "relative_path"}
            }
            for file_row in row.get("files", [])
        ]
    }


def _project_people_asset_path(root_dir: Path, project_id: str) -> Path:
    return root_dir / "project_assets" / project_id / "people_structure.json"


async def _read_upload_with_limit(
    upload: UploadFile,
    limit: int,
    *,
    total_remaining: int | None = None,
) -> bytes:
    chunks: list[bytes] = []
    total = 0
    while True:
        read_limit = min(1024 * 1024, limit - total + 1)
        if total_remaining is not None:
            read_limit = min(read_limit, total_remaining - total + 1)
        chunk = await upload.read(max(1, read_limit))
        if not chunk:
            break
        total += len(chunk)
        if total > limit:
            raise HTTPException(
                status_code=413,
                detail=(
                    f"Upload file {Path(upload.filename or 'uploaded_file').name} "
                    f"exceeds configured limit of {limit} bytes"
                ),
            )
        if total_remaining is not None and total > total_remaining:
            raise HTTPException(
                status_code=413,
                detail=(
                    f"Upload request exceeds configured total limit while reading "
                    f"{Path(upload.filename or 'uploaded_file').name}"
                ),
            )
        chunks.append(chunk)
    return b"".join(chunks)


def _persist_project_people(
    store: ProjectSQLiteStore,
    scope: Any,
    structure: Any,
    *,
    root_dir: Path,
    source_path: Path | None,
    source_type: str,
) -> Path:
    asset_path = _project_people_asset_path(root_dir, scope.project_id)
    save_people_asset(
        structure,
        asset_path,
        asset_id=f"people_structure_{scope.project_id}",
        source_type=source_type,
        source_path=str(source_path or ""),
        notes=f"Project-scoped people asset updated by {scope.actor.id}.",
    )
    project = store.get_project(scope.project_id)
    if project is None:
        raise KeyError(f"Unknown project: {scope.project_id}")
    tags = {
        "org_id": scope.actor.org_id,
        "project_id": scope.project_id,
        "topic_id": None,
        "author_id": scope.actor.id,
        "sensitivity": "l1",
    }
    raw_path = source_path or asset_path
    store.ingest(
        {
            "doc_id": f"people_structure_{scope.project_id}",
            "title": f"{project['name']} 人员结构",
            "meeting_date": date.today().isoformat(),
            "topic": "人员结构与职责分工",
            "tags": ["people_structure", source_type],
            "curated_source": {
                "source_type": "project people JSON",
                "path": str(asset_path),
                "status": "matched",
                "notes": "项目专属解析资产",
            },
            "raw_source": {
                "source_type": source_type,
                "path": str(raw_path),
                "status": "matched",
                "notes": "原始导入文件" if source_path else "人工维护项目资产",
            },
            "people_asset_path": str(asset_path),
        },
        tags=tags,
        actor=scope.actor,
        kind="people_structure",
        tag_origin="actor_ingest",
    )
    return asset_path


def _person_path(root_title: str, group: str, role: str, name: str) -> str:
    return " / ".join(item for item in [root_title, group, role, name] if item)


def _parse_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value[:10])
    except ValueError:
        return None


def _create_manual_item(
    store: ProjectSQLiteStore,
    payload: ThreeListMutationRequest,
    *,
    category: str,
    prefix: str,
    actor: Any,
) -> InspectionItem:
    now = datetime.now(timezone.utc).isoformat()
    title = (payload.title or "").strip() or "未命名条目"
    item = InspectionItem(
        item_id=f"{prefix}_{uuid4().hex[:12]}",
        category=category,
        title=title,
        description=(payload.description or title).strip(),
        evidence_refs=[_manual_evidence_ref(title)],
        status=CandidateStatus.CONFIRMED,
        owner_candidates=payload.owner_candidates or [],
        confirmation_editor=actor.id,
        confirmation_notes=payload.notes or "人工新增并确认",
        linked_issue_ids=payload.linked_issue_ids or [],
        linked_task_ids=payload.linked_task_ids or [],
        updated_at=now,
        due_date=payload.due_date,
        deliverable=payload.deliverable,
        acceptance_criteria=payload.acceptance_criteria,
        professional_id=payload.professional_id or "",
        board_id=payload.board_id or "",
        business_goal=payload.business_goal,
        principles=payload.principles,
        reasoning_chain=payload.reasoning_chain,
        applicable_scope=payload.applicable_scope,
        org_id=actor.org_id,
        project_id=store.default_project_id_for_actor(actor.id, actor.org_id),
        author_id=actor.id,
        sensitivity="l1",
        tag_origin="actor_ingest",
    )
    store.save_item(item)
    store.append_event("item", item.item_id, "manual_three_list_created", to_plain(item))
    return item


def _update_item_from_three_list(
    store: ProjectSQLiteStore,
    item_id: str,
    payload: ThreeListMutationRequest,
    *,
    editor: str,
) -> InspectionItem:
    updated = store.update_item_fields(
        item_id,
        status=payload.status,
        title=payload.title,
        description=payload.description,
        due_date=payload.due_date,
        deliverable=payload.deliverable,
        acceptance_criteria=payload.acceptance_criteria,
        professional_id=payload.professional_id,
        board_id=payload.board_id,
        owner_candidates=payload.owner_candidates,
        linked_issue_ids=payload.linked_issue_ids,
        linked_task_ids=payload.linked_task_ids,
        editor=editor,
        notes=payload.notes,
    )
    if any(value is not None for value in [
        payload.business_goal,
        payload.principles,
        payload.reasoning_chain,
        payload.applicable_scope,
    ]):
        item = store.get_item(item_id)
        if payload.business_goal is not None:
            item.business_goal = payload.business_goal
        if payload.principles is not None:
            item.principles = payload.principles
        if payload.reasoning_chain is not None:
            item.reasoning_chain = payload.reasoning_chain
        if payload.applicable_scope is not None:
            item.applicable_scope = payload.applicable_scope
        item.confirmation_editor = editor
        item.confirmation_notes = payload.notes or item.confirmation_notes
        item.updated_at = datetime.now(timezone.utc).isoformat()
        store.save_item(item)
        return item
    return updated


def _manual_evidence_ref(title: str) -> EvidenceRef:
    return EvidenceRef(
        source_doc_id="manual_three_list",
        source_kind="manual",
        locator="three_lists",
        quote=title,
        evidence_level="human_confirmed",
    )


app = create_app()
